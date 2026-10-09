# tests/documents.py
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfWriter
from docx import Document

import zipfile
import io


def make_pdf(lines: list[str], uri: str | None = None) -> bytes:
    content = (
        "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET"
    )
    annots = " /Annots [6 0 R]" if uri else ""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 5 0 R "
        f"/Resources << /Font << /F1 4 0 R >> >>{annots} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    if uri:
        objects.append(
            "<< /Type /Annot /Subtype /Link /Rect [72 700 200 716] "
            f"/A << /S /URI /URI ({uri}) >> >>"
        )
    output = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n{body}\nendobj\n".encode()
    table = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{table}\n%%EOF"
    ).encode()
    return output


def make_encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(paragraphs: list[str], table_cells: list[str] | None = None, link: str | None = None):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for cell, text in zip(table.rows[0].cells, table_cells, strict=True):
            cell.text = text
    if link:
        document.part.relate_to(link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()
